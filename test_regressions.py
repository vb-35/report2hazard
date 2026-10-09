"""Small offline research smoke check. Run with python test_regressions.py."""

import csv
import json
import os
import re
import shutil
import socket
import urllib.request
from copy import deepcopy
from dataclasses import replace
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest import TestCase
from unittest.mock import patch

from docx import Document
from pypdf import PdfReader, PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from multi_hazard_pipeline import human_review, language, pipeline, splitter, web
from multi_hazard_pipeline.agents.classification_agent import _classification_prompt, classification_agent
from multi_hazard_pipeline.agents.review_agent import CHECK_NAMES, review_agent
from multi_hazard_pipeline.agents.segment_agent import (
    CONSOLIDATION_PROMPT, SEGMENTATION_PROMPT, batch_chunks, segment_agent, _verify_nonliteral_citation, _restore_consolidated_evidence,
)
from multi_hazard_pipeline.agents.source_agent import source_agent, split_source_chunks
from multi_hazard_pipeline.agents.translation_agent import translation_agent
from multi_hazard_pipeline.config import DEFAULT_CONFIG, TAXONOMY_DECISION_RULES
from multi_hazard_pipeline.core import read_json, resolve_path, write_json
from multi_hazard_pipeline.errors import PipelineError
from multi_hazard_pipeline.llm import ChatClient
from multi_hazard_pipeline.payloads import compact_json, resolve_chunk_ids, source_payload
from multi_hazard_pipeline.schemas import CitationQuoteMismatch, validate_segment_chain


class SmokeClient:
    """Fixed model responses; exercise the real stages without an API key."""

    def __init__(self):
        self.calls = []

    def complete_json(self, **kwargs):
        name = kwargs["response_schema"]["name"]
        data = kwargs["user_payload"]
        self.calls.append((name, data))
        if name == "translated_chunks":
            payload = {"translations": [
                {"chunk_id": chunk["chunk_id"], "source_language": chunk["source_language"],
                 "translated_text": "A debris flow reached the torrent."}
                for chunk in data["chunks"]
            ]}
            bad = deepcopy(payload)
            bad["translations"][0]["chunk_id"] = "c999999"
            with TestCase().assertRaises(ValueError):
                kwargs["validate"](bad)
        elif name == "causal_segments":
            payload = {"segments": [
                {"segment": number, "causal_order": number,
                 "predecessor_segment_ids": [] if number == 1 else [number - 1],
                 "event": "Catchment storm", "process": chunk.get("translated_text", chunk["text"]),
                 "evidence": [{"chunk_id": chunk["chunk_id"], "quote": chunk["text"]}]}
                for number, chunk in enumerate(data["chunks"], start=1)
            ]}
            bad = deepcopy(payload)
            bad["segments"][0]["evidence"][0]["chunk_id"] = "c999999"
            with TestCase().assertRaises(ValueError):
                kwargs["validate"](bad)
        elif name == "classified_segments":
            payload = {"rows": [
                {"segment": item["segment"],
                 "generalized_category": "Material Mobilization" if index == 0 else "Negative Impact on permanent or temporary infrastructure",
                 "interaction_type": "Process-process" if index == 0 else "Process-structure",
                 "sediment_transport_phase": "Erosion" if index == 0 else "Dysconnectivity",
                 "classification_rationale": ["Rainfall recruits sediment." if index == 0 else "Debris blocks the bridge."]}
                for index, item in enumerate(data["segments"])
            ]}
            with TestCase().assertRaises(ValueError):
                kwargs["validate"]({"rows": payload["rows"][:-1]})
            with TestCase().assertRaises(ValueError):
                kwargs["validate"]({"rows": [payload["rows"][0]] * 2})
        elif name == "consolidated_causal_segments":
            payload = {"segments": [
                {key: item[key] for key in ("segment", "causal_order", "predecessor_segment_ids", "event", "process")}
                | {"source_segment_ids": [item["segment"]]}
                for item in data["batch_segments"]
            ], "excluded_observations": []}
        elif name == "whole_report_evaluation":
            assert len(data["candidate_report"]["rows"]) == len(data["source_chunks"])
            chunks = {chunk["chunk_id"]: chunk for chunk in data["source_chunks"]}
            for row in data["candidate_report"]["rows"]:
                for citation in row["evidence"]:
                    assert set(citation) == {"chunk_id", "quote"}
                    assert citation["quote"] in chunks[citation["chunk_id"]]["text"]
            payload = {"status": "pass", "summary": "Synthetic chain is coherent.",
                       "issues": [], "checks": dict.fromkeys(CHECK_NAMES, True)}
        else:
            raise AssertionError(f"Unexpected model call: {name}")
        kwargs["validate"](payload)
        return payload


def check_extraction_and_translation(root):
    inputs = root / "input"
    inputs.mkdir()
    path = inputs / "report.docx"
    document = Document()
    document.add_paragraph("Heavy rainfall mobilized sediment into the channel.")
    document.add_table(rows=1, cols=1).cell(0, 0).text = "The bridge became blocked by debris."
    document.save(path)
    source = source_agent([path], "report")
    assert [chunk["text"] for chunk in source["chunks"]] == [
        "Heavy rainfall mobilized sediment into the channel.", "The bridge became blocked by debris.",
    ]
    assert source["chunks"][0]["paragraph"] == 1
    assert source["chunks"][1]["table"] == source["chunks"][1]["cell"] == 1
    split = split_source_chunks(source["chunks"], 20)
    assert "".join(chunk["text"] for chunk in split) == "".join(chunk["text"] for chunk in source["chunks"])
    assert len({chunk["chunk_id"] for chunk in split}) == len(split)

    client = SmokeClient()
    translated = translation_agent(client, source)
    assert not client.calls
    assert [chunk["translated_text"] for chunk in translated["chunks"]] == [chunk["text"] for chunk in source["chunks"]]
    for original, chunk in zip(source["chunks"], translated["chunks"], strict=True):
        assert all(chunk[key] == value for key, value in original.items())

    german = deepcopy(source)
    german["chunks"] = [german["chunks"][0]]
    german["chunks"][0]["text"] = "Ein Murgang erreichte den Bach; eine Hangmure blieb am Hang."
    translated = translation_agent(client, german)
    assert translated["chunks"][0]["text"] == german["chunks"][0]["text"]
    assert translated["chunks"][0]["paragraph"] == 1
    assert translated["chunks"][0]["translated_text"] == "A debris flow reached the torrent."
    before = deepcopy(translated)
    chain = segment_agent(client, translated, DEFAULT_CONFIG)
    assert translated == before
    assert chain["segments"][0]["evidence"][0]["chunk_id"] == german["chunks"][0]["chunk_id"]
    validate_segment_chain(chain, german)
    glossary = client.calls[0][1]["terminology_mappings"]
    assert all(set(row) == {"English", "German"} for row in glossary)
    assert next(row for row in glossary if row["German"] == "Murgang")["English"] == "debris flow"
    assert language.majority_decision(100, 76, 0, .75)["decision"] == "skip_translation"
    assert language.majority_decision(100, 75, 0, .75)["decision"] == "translate"
    assert language.majority_decision(100, 70, 20, .75)["decision"] == "unresolved"
    return inputs


def check_model_payloads(root):
    chunks = [
        {"chunk_id": "permanent-document-one-page-1", "document_id": "document-one",
         "filename": "Köln.pdf", "source_type": "pdf", "page": 1, "text": "Überflutung.",
         "translated_text": "Flooding.", "source_language": "German"},
        {"chunk_id": "permanent-document-two-page-1", "document_id": "document-two",
         "filename": "Other.pdf", "source_type": "pdf", "page": 1, "text": "Sediment moved.",
         "translated_text": "Sediment moved.", "parent_chunk_id": "parent-two", "char_start": 0, "char_end": 15},
    ]
    original = deepcopy(chunks)
    projected, references = source_payload(chunks)
    assert chunks == original
    assert source_payload(chunks) == (projected, references)
    assert references == {"c1": chunks[0]["chunk_id"], "c2": chunks[1]["chunk_id"]}
    assert projected["chunks"][0]["translated_text"] == "Flooding."
    assert "translated_text" not in projected["chunks"][1]
    assert [item["document_id"] for item in projected["chunks"]] == ["d1", "d2"]
    assert [item["page"] for item in projected["chunks"]] == [1, 1]
    assert projected["documents"][0]["filename"] == "Köln.pdf"
    assert projected["chunks"][1]["char_start"] == 0
    assert projected["chunks"][1]["char_end"] == 15
    legacy = deepcopy(chunks)
    for chunk in legacy:
        del chunk["document_id"]
        chunk["doc_id"] = "one-report"
    assert len(source_payload(legacy)[0]["documents"]) == 2
    assert "Überflutung" in compact_json(projected) and "\\u" not in compact_json(projected)
    assert compact_json({"x": [1, 2]}) == '{"x":[1,2]}'
    assert batch_chunks(chunks, 36) == [chunks]  # 12 + 9 + 15; passthrough counted once.
    with TestCase().assertRaises(PipelineError):
        batch_chunks(chunks, 19)
    response = {"translations": [{"chunk_id": "c2", "translated_text": "Sediment moved."}]}
    resolved = resolve_chunk_ids(response, references, translations=True)
    assert resolved["translations"][0]["chunk_id"] == chunks[1]["chunk_id"]
    assert response["translations"][0]["chunk_id"] == "c2"
    for bad in ("c3", chunks[0]["chunk_id"]):
        with TestCase().assertRaises(ValueError):
            resolve_chunk_ids({"translations": [{"chunk_id": bad}]}, references, translations=True)

    client = SmokeClient()
    chain = segment_agent(client, {"doc_id": "identity", "chunks": chunks}, replace(DEFAULT_CONFIG, batch_max_chars=21))
    assert [row["evidence"][0]["chunk_id"] for row in chain["segments"]] == [chunk["chunk_id"] for chunk in chunks]
    assert [data["chunks"][0]["chunk_id"] for name, data in client.calls if name == "causal_segments"] == ["c1", "c1"]
    validate_segment_chain(chain, {"chunks": chunks})
    assert chunks == original

    class WireResponse:
        def __enter__(self): return self
        def __exit__(self, *args): pass
        def read(self): return b'{"ok":true}'

    def capture_wire(req, **kwargs):
        body = req.data.decode("utf-8")
        assert "Köln" in body and "\\u" not in body
        assert body == compact_json(json.loads(body))
        assert json.loads(body)["messages"][1]["content"] == compact_json(projected)
        return WireResponse()

    llm = ChatClient("offline", replace(DEFAULT_CONFIG.llm, api_base_url="https://offline.invalid"), root / "wire-timings.jsonl")
    with patch.object(urllib.request, "urlopen", capture_wire):
        assert llm.complete_json(system_prompt="Check", user_payload=projected,
                                 validate=lambda payload: None) == {"ok": True}
    assert json.loads(llm.timing_path.read_text(encoding="utf-8"))["input_chars"] == len(compact_json(projected))

    class FragmentClient:
        def complete_json(self, **kwargs):
            assert kwargs["timeout_seconds"] == 120 and kwargs["total_timeout_seconds"] == 300
            fragments = kwargs["user_payload"]["fragments"]
            assert len(fragments) == 1  # Adjacent text from a different document cannot repair this citation.
            assert fragments[0]["text"] == chunks[0]["text"] and fragments[0]["page"] == 1
            bad = {"supported": True, "fragment_ids": ["f999"], "reason": "Invented ID"}
            with TestCase().assertRaises(ValueError):
                kwargs["validate"](bad)
            result = {"supported": True, "fragment_ids": [fragments[0]["fragment_id"]], "reason": "Source supports it."}
            kwargs["validate"](result)
            return result

    repaired = _verify_nonliteral_citation(
        FragmentClient(), {"event": "Flood", "process": "Flooding"},
        CitationQuoteMismatch(1, chunks[0]["chunk_id"], "Nonliteral flooding"), chunks,
    )
    assert repaired == [{"chunk_id": chunks[0]["chunk_id"], "quote": chunks[0]["text"]}]
    assert _verify_nonliteral_citation(
        FragmentClient(), {"event": "Flood", "process": "Flooding"},
        CitationQuoteMismatch(1, legacy[0]["chunk_id"], "Nonliteral flooding"), legacy,
    ) == repaired


def check_pipeline_and_exports(root, inputs):
    client = SmokeClient()
    manifest = pipeline.run_pipeline(inputs, root / "runs", client=client)
    assert manifest["status"] == "awaiting_human_review", manifest.get("errors")
    run = Path(manifest["artifact_dir"])
    candidate = read_json(run / "candidate_report.json")
    source = read_json(run / "source.json")
    translated = read_json(run / "translated.json")
    assert source["chunks"] == [{key: value for key, value in chunk.items()
                                 if key not in {"translated_text", "source_language"}}
                                for chunk in translated["chunks"]]
    assert [name for name, _ in client.calls] == [
        "causal_segments", "classified_segments", "whole_report_evaluation",
    ]
    assert [row["process"] for row in candidate["rows"]] == [chunk["text"] for chunk in source["chunks"]]
    assert candidate["rows"][1]["evidence"][0]["provenance"]["table"] == 1
    assert [row["evidence"][0]["chunk_id"] for row in candidate["rows"]] == [chunk["chunk_id"] for chunk in source["chunks"]]
    segmentation_payload = client.calls[0][1]
    assert all("translated_text" not in chunk for chunk in segmentation_payload["chunks"])
    assert segmentation_payload["chunks"][1]["table"] == segmentation_payload["chunks"][1]["cell"] == 1
    assert "controlled_labels" not in client.calls[1][1]

    german_inputs = root / "german"
    german_inputs.mkdir()
    (german_inputs / "Murgang.txt").write_text("Ein Murgang erreichte den Bach; eine Hangmure blieb am Hang.", encoding="utf-8")
    german_run = pipeline.run_pipeline(german_inputs, root / "translated-runs", client=SmokeClient())
    assert german_run["status"] == "awaiting_human_review", german_run.get("errors")
    german_dir = Path(german_run["artifact_dir"])
    german_source = read_json(german_dir / "source.json")
    german_translated = read_json(german_dir / "translated.json")
    assert all(german_translated["chunks"][0][key] == value for key, value in german_source["chunks"][0].items())
    assert german_translated["chunks"][0]["translated_text"] == "A debris flow reached the torrent."
    assert read_json(german_dir / "candidate_report.json")["rows"][0]["evidence"][0]["chunk_id"] == german_source["chunks"][0]["chunk_id"]
    assert not (run / "final_rows.json").exists() and not (run / "final_rows.csv").exists()
    invalid = {"segments": deepcopy(candidate["rows"])}
    invalid["segments"][0]["evidence"][0]["quote"] = "Invented source evidence"
    with TestCase().assertRaises(ValueError):
        validate_segment_chain(invalid, source)

    edited = human_review.apply_candidate_edits(
        run, [{"segment": 1, "field": "interaction_type", "new_value": "Process-topography"}],
        client=client,
    )
    assert edited["status"] == "awaiting_human_review"
    candidate = read_json(run / "candidate_report.json")
    assert candidate["candidate_revision"] == 2
    assert read_json(run / "source.json") == source
    assert read_json(run / "translated.json") == translated
    assert read_json(run / "self_evaluation.json")["candidate_revision"] == 2
    history = read_json(run / "self_evaluation.json")
    write_json(run / "self_evaluation.json", history | {"candidate_revision": 1})
    with TestCase().assertRaises(PipelineError):
        human_review.approve_run(run)
    assert not (run / "final_rows.json").exists() and not (run / "final_rows.csv").exists()
    write_json(run / "self_evaluation.json", history)
    approved = human_review.approve_run(run)
    assert approved["status"] == read_json(run / "final_rows.json")["status"] == "approved"
    with (run / "final_rows.csv").open(encoding="utf-8", newline="") as handle:
        reader = csv.DictReader(handle)
        assert reader.fieldnames == list(DEFAULT_CONFIG.export_columns)
        rows = list(reader)
    assert len(rows) == 2
    assert rows[0]["process"] == source["chunks"][0]["text"]
    assert rows[0]["interaction_type"] == "Process-topography"

    failed = pipeline.run_pipeline(inputs, root / "failed", client=object())
    assert failed["status"] == "failed"
    assert failed["errors"][0]["stage"] == "segmentation"
    assert not (Path(failed["artifact_dir"]) / "final_rows.json").exists()

    manifest = pipeline.run_pipeline(inputs, root / "export_failure", client=SmokeClient())
    run = Path(manifest["artifact_dir"])
    with patch.object(human_review, "write_csv", side_effect=OSError("disk full")):
        with TestCase().assertRaises(OSError):
            human_review.approve_run(run)
    assert read_json(run / "manifest.json")["status"] == "failed"
    assert not (run / "final_rows.json").exists() and not (run / "final_rows.csv").exists()


def blocked_network(*args, **kwargs):
    raise AssertionError("The smoke check must not contact a model service or network")


def check_long_paths(root):
    base = root / "long-path-check"
    # Match the reported 220-character parent; the old temporary filename exceeds 260.
    output = base / ("x" * (220 - len(str(base.resolve())) - 1))
    resolved = resolve_path(output)
    assert resolve_path(resolved) == resolved
    if os.name == "nt":
        assert str(resolved).startswith("\\\\?\\")
        with patch.object(Path, "resolve", return_value=Path("\\\\server\\share\\reports")):
            assert str(resolve_path("unused")) == "\\\\?\\UNC\\server\\share\\reports"
    try:
        resolved.mkdir(parents=True)
        source = base / "collection.pdf"
        writer = PdfWriter()
        font = DictionaryObject({NameObject("/Type"): NameObject("/Font"),
                                 NameObject("/Subtype"): NameObject("/Type1"),
                                 NameObject("/BaseFont"): NameObject("/Helvetica")})
        texts = ["4.1 Schallerbach. Heavy rainfall mobilized sediment into the channel.",
                 "Debris blocked the bridge.",
                 "4.2 Seigesbach. Heavy rainfall mobilized sediment into the channel.",
                 "Debris blocked the bridge."]
        for text in texts:
            page = writer.add_blank_page(width=600, height=800)
            page[NameObject("/Resources")] = DictionaryObject({
                NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})})
            stream = DecodedStreamObject()
            stream.set_data(f"BT /F1 12 Tf 40 740 Td ({text}) Tj ET".encode("ascii"))
            page[NameObject("/Contents")] = stream
        writer.write(str(source))
        ranges = [splitter.EventRange(1, "4.1 Schallerbach", 1, 2),
                  splitter.EventRange(2, "4.2 Seigesbach", 3, 4)]
        assert len(str(output / ".01-4-1-schallerbach_pages-001-002-12345678.tmp")) == 268
        outputs = splitter._write_ranges(source, output, ranges)
        for index, path in enumerate(outputs):
            pages = PdfReader(str(path)).pages
            assert len(pages) == 2
            assert [page.extract_text() for page in pages] == texts[index * 2:index * 2 + 2]
        assert not list(resolved.glob("*.tmp"))
        # Force failure after one temporary PDF has been written; cleanup must still work.
        with patch.object(splitter.PdfWriter, "write", side_effect=[None, OSError("disk full")]):
            with TestCase().assertRaises(PipelineError):
                splitter._write_ranges(source, output / "failed", ranges)
        assert not list((resolved / "failed").iterdir())

        class CollectionClient(SmokeClient):
            def complete_json(self, **kwargs):
                if kwargs["response_schema"]["name"] != "event_report_separation":
                    return super().complete_json(**kwargs)
                result = {"status": "pass", "events": [
                    {"title": event.title, "start_page": event.start_page,
                     "heading_quote": event.title, "confidence": 0.99}
                    for event in ranges], "collection_end": None, "warnings": []}
                kwargs["validate"](result)
                return result

        class InlineExecutor:
            def submit(self, operation):
                operation()

        # Exercise the real web preparation, child runs, review, and download paths offline.
        artifact_root = output / ("deep-results-" * 7)
        assert len(str(artifact_root)) > 260
        app = web.create_app(artifact_root, executor=InlineExecutor())
        client = app.test_client()
        with source.open("rb") as handle:
            response = client.post("/runs", data={"files": (handle, source.name)})
        assert response.status_code == 400 and b"single report or a multi-report" in response.data
        with (output / "notes.txt").open("w", encoding="utf-8") as handle:
            handle.write("not a collection")
        with (output / "notes.txt").open("rb") as handle:
            response = client.post("/runs", data={"files": (handle, "notes.txt"), "report_mode": "multi"})
        assert response.status_code == 400 and b"exactly one PDF" in response.data
        with patch.object(web.ChatClient, "from_config", return_value=CollectionClient()), source.open("rb") as handle:
            response = client.post("/runs", data={"files": (handle, source.name), "report_mode": "multi"})
        assert response.status_code == 303
        run_id = response.headers["Location"].rsplit("/", 1)[-1]
        artifact_root = resolve_path(artifact_root)
        manifest = read_json(artifact_root / run_id / "manifest.json")
        assert manifest["status"] == "split", manifest.get("errors")
        assert len(manifest["child_runs"]) == 2
        collection = client.get(f"/runs/{run_id}/collection").get_json()
        assert collection["status"] == "awaiting_human_review"
        assert [stage["status"] for stage in collection["stages"]] == ["completed", "completed", "awaiting"]
        assert [report["status"] for report in collection["reports"]] == ["awaiting_human_review"] * 2
        page = client.get(f"/runs/{run_id}").data
        assert b"MULTI-REPORT COLLECTION" in page and b"Extracted reports" in page
        first, second = (child["run_id"] for child in manifest["child_runs"])
        assert f"/runs/{second}".encode() in client.get(f"/runs/{first}").data
        assert b"Report 2 of 2" in client.get(f"/runs/{second}").data
        assert first.encode() not in client.get("/").data
        for child in manifest["child_runs"]:
            run = artifact_root / child["run_id"]
            assert read_json(run / "manifest.json")["status"] == "awaiting_human_review"
            assert read_json(run / "manifest.json")["parent_run_id"] == run_id
            human_review.approve_run(run)
            assert len(read_json(run / "final_rows.json")["rows"]) == 2
            for route in ("", "/source/0", "/download/final_rows.csv"):
                response = client.get(f"/runs/{child['run_id']}{route}")
                assert response.status_code == 200, route
                assert response.data
                response.close()
        assert client.get("/").status_code == 200
        assert client.get(f"/runs/{run_id}/download/manifest.json").status_code == 404
        assert client.get(f"/runs/{run_id}/source/../manifest.json").status_code == 404
    finally:
        assert base.resolve().is_relative_to(root.resolve())
        shutil.rmtree(resolve_path(base), ignore_errors=False)


def check_instruction_contracts(root, inputs):
    """Check prompt wiring and structure, not the semantic behavior of a model."""
    example = SEGMENTATION_PROMPT.split("Representative Schnannerbach example:", 1)[1]
    source = {"chunks": [{"chunk_id": chunk_id, "text": text}
                         for chunk_id, text in re.findall(r'(c\d+)="([^"]+)"', example.split("Output: ", 1)[0])]}
    payload, _ = json.JSONDecoder().raw_decode(example.split("Output: ", 1)[1])
    validate_segment_chain(payload, source)
    assert all(row["predecessor_segment_ids"] == [] for row in payload["segments"])
    unsupported = deepcopy(payload)
    unsupported["segments"][1]["predecessor_segment_ids"] = [1]
    validate_segment_chain(unsupported, source)  # Earlier IDs pass structure even when the input proves no causal link.
    branch_input = re.search(r'Chronology/branch example:\s*Input: c1="([^"]+)"', SEGMENTATION_PROMPT)[1]
    branch = {"segments": [
        {"segment": number, "causal_order": number, "predecessor_segment_ids": [2] if number == 3 else [],
         "event": "Branch example", "process": sentence.strip(),
         "evidence": [{"chunk_id": "c1", "quote": sentence.strip()}]}
        for number, sentence in enumerate(branch_input.split("."), 1) if sentence.strip()
    ]}
    validate_segment_chain(branch, {"chunks": [{"chunk_id": "c1", "text": branch_input}]})
    assert len(branch["segments"]) == 4 and branch["segments"][3]["predecessor_segment_ids"] == []
    assert "Chronology, adjacency, or shared location alone does not establish causation" in CONSOLIDATION_PROMPT

    texts = [
        "Blocks destroyed the nets, releasing trapped sediment downstream.",
        "The nets trapped blocks and stopped their downstream passage.",
        "Blocks partially destroyed the nets.",
    ]
    segments = {"doc_id": "nets", "segments": [
        {"segment": number, "causal_order": number, "predecessor_segment_ids": [],
         "event": "Rockfall", "process": text, "evidence": [{"chunk_id": f"source-{number}", "quote": text}]}
        for number, text in enumerate(texts, 1)
    ]}
    labels = ["Positive Impact on permanent or temporary infrastructure",
              "Negative Impact on permanent or temporary infrastructure",
              "Impact on permanent or temporary infrastructure"]

    class TaxonomyClient:
        def complete_json(self, **kwargs):
            assert kwargs["system_prompt"].count(TAXONOMY_DECISION_RULES) == 1
            if kwargs["response_schema"]["name"] == "classified_segments":
                result = {"rows": [
                    {"segment": row["segment"], "generalized_category": labels[row["segment"] - 1],
                     "interaction_type": "Process-structure",
                     "sediment_transport_phase": "Dysconnectivity" if row["segment"] == 2 else "Transportation",
                     "classification_rationale": ["T2: " + row["evidence"][0]["quote"]]}
                    for row in kwargs["user_payload"]["segments"]
                ]}
            else:
                assert "Independently determine which labels the source supports" in kwargs["system_prompt"]
                assert "source chunk ID and short quote, and rule ID" in kwargs["system_prompt"]
                result = {"status": "pass", "summary": "Fixed fixture only.",
                          "issues": [], "checks": dict.fromkeys(CHECK_NAMES, True)}
            kwargs["validate"](result)
            return result

    client = TaxonomyClient()
    before = deepcopy(segments)
    classified = classification_agent(client, segments, DEFAULT_CONFIG)
    revised = classification_agent(
        client, segments, DEFAULT_CONFIG, segment_ids=[3], existing=classified,
        correction_instruction="Reclassify segment 3 as Negative Impact on permanent or temporary infrastructure.",
    )
    assert revised == classified and segments == before  # No automatic adoption of the proposed replacement.
    assert [row["generalized_category"] for row in revised["rows"]] == labels
    source = {"doc_id": "nets", "chunks": [{"chunk_id": f"source-{number}", "text": text}
                                            for number, text in enumerate(texts, 1)]}
    assert review_agent(client, revised, source, DEFAULT_CONFIG)["status"] == "pass"
    assert _classification_prompt(DEFAULT_CONFIG).count(TAXONOMY_DECISION_RULES) == 1

    class DisagreementClient(SmokeClient):
        def __init__(self):
            super().__init__()
            self.review_count = 0

        def complete_json(self, **kwargs):
            name = kwargs["response_schema"]["name"]
            if name == "classified_segments" and self.review_count:
                # The existing classifier call may return its supported label despite a reviewer proposal.
                assert [row["segment"] for row in kwargs["user_payload"]["segments"]] == [1, 2]
                assert kwargs["user_payload"]["requested_segment_ids"] == [2]
                assert "Positive Impact" in kwargs["user_payload"]["review_issues"][0]["suggested_action"]
                result = {"rows": [{"segment": 2, "generalized_category": labels[1],
                                    "interaction_type": "Process-structure", "sediment_transport_phase": "Dysconnectivity",
                                    "classification_rationale": ["T2: Debris blocks sediment passage at the bridge."]}]}
                kwargs["validate"](result)
                return result
            result = super().complete_json(**kwargs)
            if name == "whole_report_evaluation":
                self.review_count += 1
                if self.review_count == 1:
                    result = {"status": "revision_required", "summary": "Synthetic disagreement for correction routing.",
                              "issues": [{"stage": "categorization", "segment_ids": [2], "code": "wrong_category",
                                          "message": "Synthetic reviewer claim to reassess, not a semantic verdict.",
                                          "suggested_action": "Use Positive Impact on permanent or temporary infrastructure."}],
                              "checks": dict.fromkeys(CHECK_NAMES, True) | {"categories_valid": False}}
                    kwargs["validate"](result)
            return result

    client = DisagreementClient()
    manifest = pipeline.run_pipeline(inputs, root / "disagreement", client=client)
    assert manifest["status"] == "awaiting_human_review", manifest.get("errors")
    run = Path(manifest["artifact_dir"])
    assert read_json(run / "candidate_report.json")["rows"][1]["generalized_category"] == labels[1]
    rounds = read_json(run / "self_evaluation.json")["correction_rounds"]
    assert len(rounds) == 1 and rounds[0]["stages_rerun"] == ["categorization"] and client.review_count == 2


def check_consolidation_exclusions():
    source = {"chunks": [{"chunk_id": "c1", "text": "Rain fell. Heavy rain triggered flow. Historical flood."}]}
    preliminary = [
        {"segment": i, "causal_order": i, "predecessor_segment_ids": [], "event": "Storm",
         "process": quote, "evidence": [{"chunk_id": "c1", "quote": quote}]}
        for i, quote in enumerate(["Rain fell.", "Heavy rain triggered flow.", "Historical flood."], 1)
    ]
    payload = {"segments": [{"segment": 1, "causal_order": 1, "predecessor_segment_ids": [],
                            "event": "Storm", "process": "Heavy rainfall", "source_segment_ids": [1, 2]}],
               "excluded_observations": [{"source_segment_id": 3, "reason": "Historical event"}]}
    result = _restore_consolidated_evidence(payload, preliminary, source)
    assert result[0]["evidence"] == preliminary[0]["evidence"] + preliminary[1]["evidence"]
    assert len(result) == 1
    for exclusions in (None, [], [{"source_segment_id": 3, "reason": " "}],
                       [{"source_segment_id": 99, "reason": "Unknown"}],
                       payload["excluded_observations"] * 2,
                       payload["excluded_observations"] + [{"source_segment_id": 1, "reason": "Already retained"}]):
        invalid = deepcopy(payload)
        invalid["excluded_observations"] = exclusions
        with TestCase().assertRaises(ValueError):
            _restore_consolidated_evidence(invalid, preliminary, source)


def main():
    check_consolidation_exclusions()
    with (
        TemporaryDirectory(prefix="hazard-smoke-") as directory,
        patch.object(urllib.request.OpenerDirector, "open", blocked_network),
        patch.object(socket.socket, "connect", blocked_network),
        patch.object(socket.socket, "connect_ex", blocked_network),
    ):
        root = Path(directory)
        check_model_payloads(root)
        inputs = check_extraction_and_translation(root)
        check_pipeline_and_exports(root, inputs)
        check_long_paths(root)
        check_instruction_contracts(root, inputs)
    print("Research smoke check passed (offline; data integrity, prompt wiring, correction routing and export; no model semantic-quality claim).")


if __name__ == "__main__":
    main()
