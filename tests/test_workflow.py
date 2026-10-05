from __future__ import annotations

from copy import deepcopy
from dataclasses import replace
from concurrent.futures import ThreadPoolExecutor
import json
from pathlib import Path

import pytest
from docx import Document

from multi_hazard_pipeline import human_review, pipeline
from multi_hazard_pipeline.agents.source_agent import discover_inputs, extract_docx, source_agent
from multi_hazard_pipeline.config import DEFAULT_CONFIG
from multi_hazard_pipeline.core import read_json
from multi_hazard_pipeline.errors import PipelineError
from multi_hazard_pipeline.schemas import validate_review_payload


def test_readable_run_names(tmp_path):
    from datetime import UTC, datetime

    stamp = datetime(2026, 9, 28, 8, 36, 16, 123456, tzinfo=UTC)
    assert pipeline.new_run_id("Grünsangerlbach / Feldbach", created_at=stamp) == "grunsangerlbach-feldbach__2026-09-28_08-36-16-123456Z"
    assert pipeline.new_run_id("../", created_at=stamp).startswith("report__")
    assert len(pipeline.new_run_id("x" * 300, created_at=stamp)) < 120
    (tmp_path / "Schnannerbach.txt").write_text("Flood report.")
    first = pipeline.create_run(tmp_path, tmp_path / "results")
    second = pipeline.create_run(tmp_path, tmp_path / "results")
    assert first["run_id"].startswith("schnannerbach__")
    assert first["run_id"] != second["run_id"]


def source_payload(doc_id: str = "report") -> dict:
    return {
        "doc_id": doc_id,
        "status": "pass",
        "files": ["report.txt"],
        "chunks": [
            {
                "chunk_id": f"{doc_id}-report-txt-0001",
                "doc_id": doc_id,
                "document_id": f"{doc_id}-report-txt",
                "filename": "report.txt",
                "file": "report.txt",
                "source_type": "txt",
                "source_kind": "txt",
                "text": "Heavy rainfall mobilized sediment into the channel.",
            }
        ],
    }


def segment_payload(doc_id: str = "report") -> dict:
    return {
        "doc_id": doc_id,
        "status": "pass",
        "segments": [
            {
                "segment": 1,
                "causal_order": 1,
                "predecessor_segment_ids": [],
                "event": "Report catchment",
                "process": "Heavy rainfall mobilized sediment",
                "evidence": [
                    {
                        "chunk_id": f"{doc_id}-report-txt-0001",
                        "quote": "Heavy rainfall mobilized sediment into the channel.",
                    }
                ],
            }
        ],
    }


def classified_payload(doc_id: str = "report") -> dict:
    row = segment_payload(doc_id)["segments"][0] | {
        "generalized_category": "Material Mobilization",
        "interaction_type": "Process-process",
        "sediment_transport_phase": "Erosion",
        "classification_rationale": ["Sediment is recruited into transport."],
    }
    return {"doc_id": doc_id, "status": "pass", "rows": [row]}


def evaluation(status: str = "pass", stage: str = "categorization") -> dict:
    issues = []
    checks = {
        "causal_chain_coherent": True,
        "evidence_verified": True,
        "segments_complete": True,
        "segments_unique": True,
        "categories_valid": True,
    }
    if status == "revision_required":
        issues = [
            {
                "stage": stage,
                "segment_ids": [1],
                "code": f"test_{stage}_issue",
                "message": f"Fix {stage}",
                "suggested_action": f"Rerun {stage}",
            }
        ]
        checks[
            "evidence_verified"
            if stage == "translation"
            else "causal_chain_coherent"
            if stage == "segmentation"
            else "categories_valid"
        ] = False
    return {"status": status, "summary": status, "issues": issues, "checks": checks}


def input_directory(tmp_path: Path) -> Path:
    input_dir = tmp_path / "report"
    input_dir.mkdir()
    (input_dir / "report.txt").write_text(
        "Heavy rainfall mobilized sediment into the channel.", encoding="utf-8"
    )
    return input_dir


def install_stage_fakes(monkeypatch: pytest.MonkeyPatch, evaluations: list[dict]) -> dict[str, int]:
    calls = {"translation": 0, "segmentation": 0, "categorization": 0, "evaluation": 0}

    def fake_translation(client, source, correction_instruction=None):
        calls["translation"] += 1
        translated = deepcopy(source)
        for chunk in translated["chunks"]:
            chunk["source_language"] = "English"
            chunk["translated_text"] = chunk["text"]
        return translated

    def fake_segments(client, source, config, correction_instruction=None):
        calls["segmentation"] += 1
        payload = segment_payload(source["doc_id"])
        payload["segments"][0]["evidence"][0] = {
            "chunk_id": source["chunks"][0]["chunk_id"],
            "quote": source["chunks"][0]["text"],
        }
        return payload

    def fake_classification(
        client,
        segments,
        config,
        segment_ids=None,
        existing=None,
        correction_instruction=None,
    ):
        calls["categorization"] += 1
        assert segment_ids is None or segment_ids == [1]
        row = dict(segments["segments"][0])
        row.update(
            {
                "generalized_category": "Material Mobilization",
                "interaction_type": "Process-process",
                "sediment_transport_phase": "Erosion",
                "classification_rationale": ["Sediment is recruited into transport."],
            }
        )
        return {"doc_id": segments["doc_id"], "status": "pass", "rows": [row]}

    queue = list(evaluations)

    def fake_evaluation(client, candidate, source, config):
        calls["evaluation"] += 1
        return queue.pop(0) if queue else evaluations[-1]

    monkeypatch.setattr(pipeline, "translation_agent", fake_translation)
    monkeypatch.setattr(pipeline, "segment_agent", fake_segments)
    monkeypatch.setattr(pipeline, "classification_agent", fake_classification)
    monkeypatch.setattr(pipeline, "review_agent", fake_evaluation)
    return calls


def run_with_fakes(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    evaluations: list[dict],
    max_rounds: int = 3,
) -> tuple[dict, dict[str, int]]:
    calls = install_stage_fakes(monkeypatch, evaluations)
    config = replace(DEFAULT_CONFIG, max_correction_rounds=max_rounds)
    manifest = pipeline.run_pipeline(
        input_directory(tmp_path), tmp_path / "runs", config, client=object()
    )
    return manifest, calls


def test_run_writes_per_stage_timing_sidecar(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    manifest, _ = run_with_fakes(tmp_path, monkeypatch, [evaluation()])
    records = [json.loads(line) for line in (Path(manifest["artifact_dir"]) / "timings.jsonl").read_text(encoding="utf-8").splitlines()]
    assert [record["stage"] for record in records if record["type"] == "stage"] == [
        "extraction", "translation", "segmentation", "categorization", "candidate_report", "self_evaluation",
    ]
    assert all(record["elapsed_seconds"] >= 0 for record in records)


def test_docx_extraction_preserves_paragraph_provenance(tmp_path: Path) -> None:
    path = tmp_path / "sample.docx"
    document = Document()
    document.add_paragraph("")
    document.add_paragraph("  Flood   mobilized sediment.  ")
    document.save(path)
    chunks, next_id = extract_docx(path, "report", 1)
    assert next_id == 2
    assert chunks[0]["paragraph"] == 2
    assert chunks[0]["text"] == "Flood mobilized sediment."
    assert chunks[0]["filename"] == "sample.docx"
    assert chunks[0]["source_type"] == "docx"


def test_input_directory_must_exist_and_contain_supported_files(tmp_path: Path) -> None:
    with pytest.raises(PipelineError, match="does not exist"):
        discover_inputs(tmp_path / "missing")
    empty = tmp_path / "empty"
    empty.mkdir()
    with pytest.raises(PipelineError, match="no supported"):
        discover_inputs(empty)


def test_chunk_ids_stay_stable_when_an_earlier_file_is_added(tmp_path: Path) -> None:
    report_dir = tmp_path / "report"
    report_dir.mkdir()
    later = report_dir / "b.txt"
    later.write_text("Later file text", encoding="utf-8")
    first_id = source_agent([later], "report")["chunks"][0]["chunk_id"]
    earlier = report_dir / "a.txt"
    earlier.write_text("Earlier file text", encoding="utf-8")
    second_id = next(
        chunk["chunk_id"]
        for chunk in source_agent([earlier, later], "report")["chunks"]
        if chunk["filename"] == "b.txt"
    )
    assert first_id == second_id


def test_categorization_issue_reruns_only_categorization(tmp_path: Path, monkeypatch) -> None:
    manifest, calls = run_with_fakes(
        tmp_path, monkeypatch, [evaluation("revision_required", "categorization"), evaluation()]
    )
    assert manifest["status"] == "awaiting_human_review"
    assert calls == {"translation": 1, "segmentation": 1, "categorization": 2, "evaluation": 2}
    history = read_json(Path(manifest["artifact_dir"]) / "self_evaluation.json")
    assert history["correction_rounds"][0]["stages_rerun"] == ["categorization"]


def test_segmentation_issue_reruns_dependent_stages(tmp_path: Path, monkeypatch) -> None:
    manifest, calls = run_with_fakes(
        tmp_path, monkeypatch, [evaluation("revision_required", "segmentation"), evaluation()]
    )
    assert manifest["status"] == "awaiting_human_review"
    assert calls == {"translation": 1, "segmentation": 2, "categorization": 2, "evaluation": 2}
    history = read_json(Path(manifest["artifact_dir"]) / "self_evaluation.json")
    assert history["correction_rounds"][0]["stages_rerun"] == [
        "segmentation",
        "categorization",
    ]


def test_every_correction_reruns_whole_report_evaluation(tmp_path: Path, monkeypatch) -> None:
    manifest, calls = run_with_fakes(
        tmp_path,
        monkeypatch,
        [
            evaluation("revision_required", "categorization"),
            evaluation("revision_required", "categorization"),
            evaluation(),
        ],
    )
    assert manifest["correction_rounds"] == 2
    assert calls["evaluation"] == 3


def test_translation_issue_reruns_translation_and_dependent_stages(tmp_path: Path, monkeypatch) -> None:
    manifest, calls = run_with_fakes(
        tmp_path, monkeypatch, [evaluation("revision_required", "translation"), evaluation()]
    )
    assert manifest["status"] == "awaiting_human_review"
    assert calls == {"translation": 2, "segmentation": 2, "categorization": 2, "evaluation": 2}
    history = read_json(Path(manifest["artifact_dir"]) / "self_evaluation.json")
    assert history["correction_rounds"][0]["stages_rerun"] == [
        "translation",
        "segmentation",
        "categorization",
    ]


def test_correction_stops_at_configured_maximum(tmp_path: Path, monkeypatch) -> None:
    manifest, calls = run_with_fakes(
        tmp_path, monkeypatch, [evaluation("revision_required", "categorization")], max_rounds=2
    )
    assert manifest["status"] == "revision_required"
    assert manifest["correction_rounds"] == 2
    assert calls == {"translation": 1, "segmentation": 1, "categorization": 3, "evaluation": 3}


def test_passing_evaluation_waits_for_human_without_final_exports(tmp_path: Path, monkeypatch) -> None:
    manifest, _ = run_with_fakes(tmp_path, monkeypatch, [evaluation()])
    run_dir = Path(manifest["artifact_dir"])
    assert manifest["status"] == "awaiting_human_review"
    assert (run_dir / "human_review.json").is_file()
    assert (run_dir / "translated.json").is_file()
    assert manifest["artifacts"]["translated"].endswith("translated.json")
    assert not (run_dir / "final_rows.json").exists()
    assert not (run_dir / "final_rows.csv").exists()


def test_human_approval_creates_authoritative_exports(tmp_path: Path, monkeypatch) -> None:
    manifest, _ = run_with_fakes(tmp_path, monkeypatch, [evaluation()])
    approved = human_review.approve_run(manifest["artifact_dir"], global_comment="Looks correct")
    run_dir = Path(approved["artifact_dir"])
    assert approved["status"] == "approved"
    assert read_json(run_dir / "final_rows.json")["status"] == "approved"
    assert (run_dir / "final_rows.csv").is_file()


def test_human_decisions_are_serialized_per_run(tmp_path: Path, monkeypatch) -> None:
    manifest, _ = run_with_fakes(tmp_path, monkeypatch, [evaluation()])
    run_dir = Path(manifest["artifact_dir"])

    def decide(action):
        try:
            return action(run_dir)["status"]
        except PipelineError:
            return "conflict"

    with ThreadPoolExecutor(max_workers=2) as executor:
        results = list(executor.map(decide, (human_review.approve_run, human_review.reject_run)))
    final_status = read_json(run_dir / "manifest.json")["status"]
    assert results.count("conflict") == 1
    assert final_status in {"approved", "rejected"}
    assert (run_dir / "final_rows.json").exists() is (final_status == "approved")
    assert (run_dir / "final_rows.csv").exists() is (final_status == "approved")


def test_failed_approval_leaves_no_authoritative_export(tmp_path: Path, monkeypatch) -> None:
    manifest, _ = run_with_fakes(tmp_path, monkeypatch, [evaluation()])
    run_dir = Path(manifest["artifact_dir"])

    def fail_csv(*args, **kwargs):
        raise OSError("disk full")

    monkeypatch.setattr(human_review, "write_csv", fail_csv)
    with pytest.raises(OSError, match="disk full"):
        human_review.approve_run(run_dir)
    failed = read_json(run_dir / "manifest.json")
    assert failed["status"] == "failed"
    assert failed["errors"][-1]["stage"] == "final_export"
    assert not (run_dir / "final_rows.json").exists()
    assert not (run_dir / "final_rows.csv").exists()


def test_human_rejection_creates_no_approved_export(tmp_path: Path, monkeypatch) -> None:
    manifest, _ = run_with_fakes(tmp_path, monkeypatch, [evaluation()])
    rejected = human_review.reject_run(manifest["artifact_dir"], global_comment="Unsupported")
    run_dir = Path(rejected["artifact_dir"])
    assert rejected["status"] == "rejected"
    assert not (run_dir / "final_rows.json").exists()
    assert not (run_dir / "final_rows.csv").exists()


def test_approval_is_blocked_without_passing_self_evaluation(tmp_path: Path, monkeypatch) -> None:
    manifest, _ = run_with_fakes(
        tmp_path, monkeypatch, [evaluation("revision_required")], max_rounds=0
    )
    with pytest.raises(PipelineError, match="self-evaluation-passed"):
        human_review.approve_run(manifest["artifact_dir"])


def test_human_edits_are_audited_and_reevaluated(tmp_path: Path, monkeypatch) -> None:
    manifest, _ = run_with_fakes(tmp_path, monkeypatch, [evaluation()])
    reevaluations = []

    def fake_review(client, candidate, source, config):
        reevaluations.append(candidate["candidate_revision"])
        return evaluation()

    monkeypatch.setattr(human_review, "review_agent", fake_review)
    changed = human_review.apply_candidate_edits(
        manifest["artifact_dir"],
        [{"segment": 1, "field": "interaction_type", "new_value": "Process-topography"}],
        client=object(),
    )
    audit = read_json(Path(changed["artifact_dir"]) / "human_review.json")["edits"]
    assert changed["status"] == "awaiting_human_review"
    assert reevaluations == [2]
    assert audit[0]["old_value"] == "Process-process"
    assert audit[0]["new_value"] == "Process-topography"
    assert changed["stages"]["self_evaluation"] == "completed"
    assert changed["progress"] == 90


def test_blank_human_segment_text_is_rejected_deterministically(tmp_path: Path, monkeypatch) -> None:
    manifest, _ = run_with_fakes(tmp_path, monkeypatch, [evaluation()])
    with pytest.raises(PipelineError, match="has no process"):
        human_review.apply_candidate_edits(
            manifest["artifact_dir"],
            [{"segment": 1, "field": "process", "new_value": "  "}],
            client=object(),
        )


def test_failed_edit_evaluation_retains_audit(tmp_path: Path, monkeypatch) -> None:
    manifest, _ = run_with_fakes(tmp_path, monkeypatch, [evaluation()])

    def fail_review(*args, **kwargs):
        raise RuntimeError("evaluation unavailable")

    monkeypatch.setattr(human_review, "review_agent", fail_review)
    with pytest.raises(RuntimeError, match="evaluation unavailable"):
        human_review.apply_candidate_edits(
            manifest["artifact_dir"],
            [{"segment": 1, "field": "interaction_type", "new_value": "Process-topography"}],
            client=object(),
        )
    human = read_json(Path(manifest["artifact_dir"]) / "human_review.json")
    assert human["edits"][-1]["old_value"] == "Process-process"
    assert human["edits"][-1]["new_value"] == "Process-topography"


def test_human_requested_correction_routes_to_requested_stage(tmp_path: Path, monkeypatch) -> None:
    manifest, _ = run_with_fakes(tmp_path, monkeypatch, [evaluation()])
    captured = {}

    def fake_corrections(**kwargs):
        captured["issue"] = kwargs["history"]["latest_evaluation"]["issues"][0]
        candidate = kwargs["candidate"]
        candidate["candidate_revision"] += 1
        kwargs["history"]["status"] = "pass"
        kwargs["history"]["latest_evaluation"] = evaluation()
        return kwargs["segments"], kwargs["classified"], candidate, kwargs["history"]

    monkeypatch.setattr(human_review, "correct_until_terminal", fake_corrections)
    corrected = human_review.request_correction(
        manifest["artifact_dir"],
        requested_stage="categorization",
        segment_ids=[1],
        segment_comments=[
            {"segment": 1, "issue_type": "miscategorized", "comment": "Use mobilization"}
        ],
        client=object(),
    )
    assert corrected["status"] == "awaiting_human_review"
    assert captured["issue"]["stage"] == "categorization"
    assert captured["issue"]["segment_ids"] == [1]
    assert "Use mobilization" in captured["issue"]["suggested_action"]
    validate_review_payload(
        {
            "status": "revision_required",
            "summary": captured["issue"]["message"],
            "issues": [captured["issue"]],
            "checks": {
                "causal_chain_coherent": True,
                "evidence_verified": True,
                "segments_complete": True,
                "segments_unique": True,
                "categories_valid": False,
            },
        }
    )


def test_human_correction_respects_run_wide_maximum(tmp_path: Path, monkeypatch) -> None:
    manifest, _ = run_with_fakes(tmp_path, monkeypatch, [evaluation()], max_rounds=0)
    with pytest.raises(PipelineError, match="maximum correction rounds"):
        human_review.request_correction(
            manifest["artifact_dir"],
            requested_stage="categorization",
            segment_ids=[1],
            client=object(),
        )


def test_unexpected_failure_retains_structured_diagnostics(tmp_path: Path, monkeypatch) -> None:
    def explode(*args, **kwargs):
        raise RuntimeError("model unavailable")

    monkeypatch.setattr(pipeline, "segment_agent", explode)
    monkeypatch.setattr(
        pipeline,
        "translation_agent",
        lambda client, source: source
        | {
            "chunks": [
                chunk | {"source_language": "English", "translated_text": chunk["text"]}
                for chunk in source["chunks"]
            ]
        },
    )
    manifest = pipeline.run_pipeline(
        input_directory(tmp_path), tmp_path / "runs", client=object()
    )
    assert manifest["status"] == "failed"
    assert manifest["errors"][0]["stage"] == "segmentation"
    assert manifest["errors"][0]["type"] == "RuntimeError"
    assert "model unavailable" in manifest["errors"][0]["message"]


def test_docx_tables_and_nested_cells_preserve_order_and_provenance(tmp_path):
    path = tmp_path / "tables.docx"
    document = Document()
    document.add_paragraph("Before the flood.")
    table = document.add_table(rows=1, cols=2)
    cell = table.cell(0, 0).merge(table.cell(0, 1))
    cell.text = "The bridge collapsed."
    cell.add_table(rows=1, cols=1).cell(0, 0).text = "Sediment was released."
    document.add_paragraph("After the flood.")
    document.save(path)
    source = source_agent([path], "report")
    assert [chunk["text"] for chunk in source["chunks"]] == [
        "Before the flood.", "The bridge collapsed.", "Sediment was released.", "After the flood.",
    ]
    assert len({chunk["chunk_id"] for chunk in source["chunks"]}) == 4
    chunk = source["chunks"][1]
    assert (chunk["table"], chunk["row"], chunk["cell"]) == (1, 1, 1)
    assert source["chunks"][-1]["paragraph"] == 2
    classified = classified_payload()
    classified["rows"][0]["evidence"] = [{"chunk_id": chunk["chunk_id"], "quote": chunk["text"]}]
    candidate = pipeline.build_candidate_report("run", classified, source, 1)
    assert candidate["rows"][0]["evidence"][0]["provenance"]["table_path"] == chunk["table_path"]


@pytest.mark.parametrize("status", ["pass", "revision_required"])
def test_initial_human_review_matches_manifest_and_evaluation(tmp_path, monkeypatch, status):
    manifest, _ = run_with_fakes(tmp_path, monkeypatch, [evaluation(status)], max_rounds=0)
    human = read_json(Path(manifest["artifact_dir"]) / "human_review.json")
    assert human["status"] == manifest["status"]
    assert human["self_evaluation_status"] == status
